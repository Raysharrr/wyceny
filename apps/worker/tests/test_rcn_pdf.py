"""/rcn-pdf-to-xlsx (T-22). Offline; the PDF is written in the test, all data invented."""

import base64
import hashlib
import hmac
import json
import time
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import main, rcn_pdf
from tests.test_rcn_pdf_words import pdf

client = TestClient(main.app)
SECRET = "test-secret"

PRINTOUT = pdf(
    [
        (399, 61, "WYDRUK Z RCN"),
        (31, 119, "GKG.GZW.4061.0000.2026"),
        (31, 210, "1"),
        (98, 210, "AAAA0000BBBB1"),
        (334, 210, "2026-07-27"),
        (31, 250, "wolny rynek"),
        (98, 250, "870 000.00"),
        (31, 300, "Lokal"),
        (419, 300, "3. mieszkalna"),
        (541, 300, "3. 88.72"),
        (734, 300, "3. ul. Lipowa 4 m.1, Testowo"),
    ]
)


def mint(exp_offset: int = 300) -> str:
    exp = int(time.time()) + exp_offset
    sig = hmac.new(SECRET.encode(), f"{exp}.cafe0123".encode(), hashlib.sha256).hexdigest()
    return f"{exp}.cafe0123.{sig}"


@pytest.fixture(autouse=True)
def secret_env(monkeypatch):
    monkeypatch.setenv("WORKER_SHARED_SECRET", SECRET)


def post(token: str, content: bytes = PRINTOUT, name="wydruk.pdf", mime="application/pdf"):
    return client.post(
        "/rcn-pdf-to-xlsx", data={"token": token}, files={"file": (name, content, mime)}
    )


def test_printout_comes_back_as_counters_and_a_workbook():
    r = post(mint())
    assert r.status_code == 200
    body = r.json()
    assert set(body) == {
        "orderNumber",
        "unit",
        "count",
        "flaggedRows",
        "fileWarnings",
        "xlsxBase64",
    }
    assert (body["orderNumber"], body["count"], body["flaggedRows"], body["fileWarnings"]) == (
        "GKG.GZW.4061.0000.2026",
        1,
        0,
        [],
    )
    ws = load_workbook(BytesIO(base64.b64decode(body["xlsxBase64"])))["transakcje"]
    assert ws["G3"].value == 870000 and ws["H3"].value == "=G3/F3"


def test_bad_token_is_401():
    assert post("nope").status_code == 401


def test_not_a_pdf_is_415():
    assert post(mint(), b"PK", "rejestr.xlsx", "application/zip").status_code == 415


def test_too_big_is_413(monkeypatch):
    monkeypatch.setattr(rcn_pdf, "MAX_PDF_BYTES", 10)
    assert post(mint()).status_code == 413


@pytest.mark.parametrize(
    ("content", "code"),
    [
        (pdf([]), "no_text_layer"),
        (pdf([(31, 60, "Umowa sprzedazy")]), "not_rcn_printout"),
        (pdf([(399, 61, "WYDRUK Z RCN")]), "no_transactions"),
    ],
)
def test_refusals_are_422_with_a_code(content, code):
    r = post(mint(), content)
    assert (r.status_code, r.json()["detail"]) == (422, {"code": code})


def test_garbage_bytes_are_422():
    r = post(mint(), b"%PDF-1.4 not really")
    assert (r.status_code, r.json()["detail"]) == (422, {"code": "not_rcn_printout"})


# R1/R4: the success log is the office's only trace of a conversion, so it must
# mean "the file is on its way back", and it must carry counters and nothing else.

STRUCTLOG_ALWAYS = {"event", "level", "timestamp", "trace_id"}


def events(captured: str, name: str) -> list[dict]:
    lines = [json.loads(line) for line in captured.splitlines() if line.strip()]
    assert lines, "no log line captured — the test would be blind"
    return [line for line in lines if line["event"] == name]


def test_success_is_logged_with_counters_only(capsys):
    r = post(mint())
    assert r.status_code == 200
    (done,) = events(capsys.readouterr().out, "rcn_pdf_converted")
    assert set(done) - STRUCTLOG_ALWAYS == {"pages", "rows", "warnings", "ms"}
    assert (done["pages"], done["rows"], done["warnings"]) == (1, 1, 0)


def test_a_workbook_that_cannot_be_built_is_not_logged_as_converted(monkeypatch, capsys):
    def boom(_printout):
        raise RuntimeError("openpyxl said no")

    monkeypatch.setattr(main.rcn_xlsx, "to_xlsx", boom)
    silent = TestClient(main.app, raise_server_exceptions=False)
    r = silent.post(
        "/rcn-pdf-to-xlsx",
        data={"token": mint()},
        files={"file": ("wydruk.pdf", PRINTOUT, "application/pdf")},
    )
    assert r.status_code == 500
    assert events(capsys.readouterr().out, "rcn_pdf_converted") == []

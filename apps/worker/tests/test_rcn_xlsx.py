"""rcn_xlsx.to_xlsx — a sheet per kind: title, headers, typed cells, a live formula, an empty J."""

from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.rcn_pdf import KIND_BUILT, KIND_PLOT, Printout, Row
from app.rcn_xlsx import to_xlsx

CLEAN = Row(1, "2026-07-27", "Testowo", "Lipowa", "4", "1", 88.72, 870000.0, False)
FLAGGED = Row(
    2, "2026-06-17", "Testowo", "Polna", "21E", "", None, 412000.0, False, ["address_unparsed"]
)
BUILT = Row(
    3,
    "2026-05-18",
    "Testowo",
    "Lipowa",
    "6",
    "",
    None,
    1280000.0,
    False,
    ["many_plots"],
    kind=KIND_BUILT,
    plot_ids=["000000_0.0001.12/3", "000000_0.0001.12/4"],
    plot_area_m2=984,
    share="1/2",
)
PLOT = Row(
    4,
    "2026-05-29",
    "Wzorcowo",
    "",
    "",
    "",
    None,
    5300000.0,
    False,
    [],
    kind=KIND_PLOT,
    plot_ids=["000000_0.0001.1043"],
    plot_area_m2=26485,
    plan="budownictwo mieszkaniowe wielorodzinne",
)


def book(*rows: Row):
    data = to_xlsx(Printout("GKG.GZW.4061.0000.2026", "000000_0 - Przykładowo", list(rows), []))
    return load_workbook(BytesIO(data))


def sheet(*rows: Row, name: str = "lokale"):
    return book(*rows)[name]


def test_title_and_headers():
    ws = sheet(CLEAN)
    assert ws["A1"].value == "WYDRUK Z RCN · GKG.GZW.4061.0000.2026 · 000000_0 - Przykładowo"
    assert [c.value for c in ws[2]] == [
        "DATA",
        "MIEJSCOWOŚĆ",
        "ULICA",
        "NR BUD",
        "NR LOK",
        "PU",
        "CENA",
        "CENA J",
        "P.P",
        "RODZAJ BUD",
    ]


def test_cells_are_typed_and_price_per_metre_is_a_formula():
    ws = sheet(CLEAN)
    assert [c.value for c in ws[3]] == [
        datetime(2026, 7, 27),
        "Testowo",
        "Lipowa",
        4,
        1,
        88.72,
        870000,
        "=G3/F3",
        "nie",
        None,
    ]
    assert ws["A3"].number_format == "DD.MM.YYYY"


def test_flagged_row_is_yellow_commented_and_has_no_formula():
    ws = sheet(CLEAN, FLAGGED)
    assert ws["D4"].value == "21E"
    assert ws["H4"].value is None
    assert ws["A4"].fill.fgColor.rgb.endswith("FFF2CC")
    assert "rozbić adresu" in ws["A4"].comment.text
    assert ws["A3"].comment is None


def test_the_empty_column_says_why_it_is_empty():
    assert "Do uzupełnienia ręcznie" in sheet(CLEAN)["J2"].comment.text


# --- domy i działki -----------------------------------------------------------


def test_every_kind_present_gets_its_own_sheet_in_order():
    assert book(PLOT, BUILT, CLEAN).sheetnames == ["lokale", "zabudowane", "niezabudowane"]


def test_a_printout_of_bare_land_alone_is_one_sheet():
    assert book(PLOT).sheetnames == ["niezabudowane"]


def test_built_sheet_columns_and_cells():
    ws = sheet(BUILT, name="zabudowane")
    assert [c.value for c in ws[2]] == [
        "DATA",
        "MIEJSCOWOŚĆ",
        "ULICA",
        "NR BUD",
        "ID DZIAŁKI",
        "POW GR",
        "UDZIAŁ",
        "CENA",
        "CENA J",
    ]
    assert [c.value for c in ws[3]] == [
        datetime(2026, 5, 18),
        "Testowo",
        "Lipowa",
        6,
        "000000_0.0001.12/3; 000000_0.0001.12/4",
        984,
        "1/2",
        1280000,
        "=H3/F3",
    ]
    assert "kilka działek" in ws["A3"].comment.text


def test_bare_land_sheet_columns_and_cells():
    ws = sheet(PLOT, name="niezabudowane")
    assert [c.value for c in ws[2]] == [
        "DATA",
        "MIEJSCOWOŚĆ",
        "ID DZIAŁKI",
        "POW GR",
        "UDZIAŁ",
        "PRZEZNACZENIE",
        "CENA",
        "CENA J",
    ]
    assert [c.value for c in ws[3]] == [
        datetime(2026, 5, 29),
        "Wzorcowo",
        "000000_0.0001.1043",
        26485,
        None,  # a whole right says nothing, and openpyxl stores "" as a blank cell
        "budownictwo mieszkaniowe wielorodzinne",
        5300000,
        "=G3/D3",
    ]
    assert ws["A3"].comment is None


def test_land_without_an_area_has_no_unit_price():
    empty = Row(5, "2026-05-29", "Wzorcowo", "", "", "", None, 1.0, False, ["missing_field"])
    empty.kind, empty.plot_area_m2 = KIND_PLOT, None
    ws = sheet(empty, name="niezabudowane")
    assert (ws["D3"].value, ws["H3"].value) == (None, None)
    assert "brakuje daty" in ws["A3"].comment.text


def test_a_partial_share_is_explained_in_the_comment():
    row = Row(6, "2026-05-29", "Wzorcowo", "", "", "", None, 1.0, False, ["partial_share"])
    row.kind, row.plot_ids, row.plot_area_m2, row.share = KIND_PLOT, ["a"], 10, "3/18"
    assert "Udział w prawie" in sheet(row, name="niezabudowane")["A3"].comment.text

"""Raw cell reader for cooperative transaction registers (T-13, S2a).

Returns the sheets of an XLSX as TEXT cells — no interpretation, no column
guessing: four known registers have five column layouts, so a human maps the
columns in the web app and `domain/coop-import.ts` does the normalising.
The only conversions here are representational: dates → ISO, numbers → their
plain string, empty → "". Rows are padded to the sheet's widest row so the
web side can index columns without bounds checks.

openpyxl `read_only=True` streams rows (the largest known register is 9.1 MB
with five sheets) and never touches formulas: `data_only=True` yields the
cached value the author saw.
"""

from __future__ import annotations

import io
from datetime import date, datetime

import openpyxl

MAX_XLSX_BYTES = 12 * 1024 * 1024  # largest known register: 9.1 MB, 5 sheets

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class NotAWorkbook(ValueError):
    """The bytes are not an XLSX workbook (XLS/CSV/PDF/scan) — not retryable."""


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        # A date cell with no time part is a date; anything else keeps the time.
        return (
            value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_sheets(data: bytes) -> list[dict]:
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a zoo: BadZipFile, InvalidFileException, KeyError…
        raise NotAWorkbook(type(exc).__name__) from exc
    sheets = []
    for ws in workbook.worksheets:
        rows = [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
        cols = max((len(r) for r in rows), default=0)
        sheets.append(
            {"name": ws.title, "cols": cols, "rows": [r + [""] * (cols - len(r)) for r in rows]}
        )
    workbook.close()
    return sheets

"""Printout rows -> the office's XLSX layout (T-22). Pure — no FastAPI import.

One sheet per kind PRESENT in the printout, in a fixed order. A flat, a house
and a bare plot answer different questions, so they get different columns rather
than one table with holes in it.
"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

from app.rcn_pdf import KIND_BUILT, KIND_PLOT, KIND_UNIT, Printout, Row

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = {
    KIND_UNIT: (
        "DATA",
        "MIEJSCOWOŚĆ",
        "ULICA",
        "NR BUD",
        "NR LOK",
        "PU",
        "CENA",
        "CENA J",
        "P.P",
        "RODZAJ BUD",
    ),
    KIND_BUILT: (
        "DATA",
        "MIEJSCOWOŚĆ",
        "ULICA",
        "NR BUD",
        "ID DZIAŁKI",
        "POW GR",
        "UDZIAŁ",
        "CENA",
        "CENA J",
    ),
    KIND_PLOT: (
        "DATA",
        "MIEJSCOWOŚĆ",
        "ID DZIAŁKI",
        "POW GR",
        "UDZIAŁ",
        "PRZEZNACZENIE",
        "CENA",
        "CENA J",
    ),
}
SHEETS = {KIND_UNIT: "lokale", KIND_BUILT: "zabudowane", KIND_PLOT: "niezabudowane"}
WIDTHS = {
    KIND_UNIT: (12, 18, 26, 8, 8, 9, 13, 11, 6, 16),
    KIND_BUILT: (12, 18, 26, 8, 30, 10, 10, 13, 11),
    KIND_PLOT: (12, 18, 30, 10, 10, 44, 13, 11),
}
WARNING_FILL = PatternFill("solid", fgColor="FFF2CC")
WARNING_TEXT = {
    "many_units": "Transakcja obejmuje kilka lokali — wpisano pierwszy mieszkalny.",
    "many_plots": "Transakcja obejmuje kilka działek — POW GR to ich suma.",
    "partial_share": "Udział w prawie inny niż 1/1 — sprawdź, czy cena dotyczy całości.",
    "address_unparsed": "Nie udało się rozbić adresu — cały adres jest w kolumnie ULICA.",
    "missing_field": "W wydruku brakuje daty, ceny albo powierzchni.",
}


def _cell_number(text: str):
    return int(text) if text.isdigit() else text or None


def _values(row: Row, n: int) -> list:
    """One transaction as the cells of its own sheet. CENA J is a live formula,
    so a corrected area or price recomputes itself in Excel."""
    when = date.fromisoformat(row.date) if row.date else None
    plots = "; ".join(row.plot_ids)
    if row.kind == KIND_UNIT:
        return [
            when,
            row.town,
            row.street,
            _cell_number(row.building),
            _cell_number(row.unit),
            row.area,
            row.price,
            f"=G{n}/F{n}" if row.area else None,
            "tak" if row.annex else "nie",
            None,
        ]
    if row.kind == KIND_BUILT:
        return [
            when,
            row.town,
            row.street,
            _cell_number(row.building),
            plots,
            row.plot_area_m2,
            row.share,
            row.price,
            f"=H{n}/F{n}" if row.plot_area_m2 else None,
        ]
    return [
        when,
        row.town,
        plots,
        row.plot_area_m2,
        row.share,
        row.plan,
        row.price,
        f"=G{n}/D{n}" if row.plot_area_m2 else None,
    ]


def to_xlsx(printout: Printout) -> bytes:
    book = Workbook()
    # openpyxl hands over a live default sheet; a printout with no flats in it
    # must not ship a stray empty "Sheet" next to the real ones.
    book.remove(book.active)
    title = " · ".join(p for p in ("WYDRUK Z RCN", printout.order_number, printout.unit) if p)
    for kind, name in SHEETS.items():
        rows = [r for r in printout.rows if r.kind == kind]
        if not rows:
            continue
        headers = HEADERS[kind]
        sheet = book.create_sheet(name)
        sheet.append([title])
        sheet.append(list(headers))
        for cell in sheet[2]:
            cell.font = Font(bold=True)
        if kind == KIND_UNIT:
            sheet.cell(2, 10).comment = Comment(
                "Do uzupełnienia ręcznie — wydruk podaje tylko „Mieszkalny”, bez rodzaju zabudowy.",
                "Wyceny",
            )
        for n, row in enumerate(rows, start=3):
            sheet.append(_values(row, n))
            sheet.cell(n, 1).number_format = "DD.MM.YYYY"
            if row.warnings:
                for cell in sheet[n]:
                    cell.fill = WARNING_FILL
                sheet.cell(n, 1).comment = Comment(
                    " ".join(WARNING_TEXT[w] for w in row.warnings), "Wyceny"
                )
        for column, width in zip((c.column_letter for c in sheet[2]), WIDTHS[kind], strict=True):
            sheet.column_dimensions[column].width = width
    out = BytesIO()
    book.save(out)
    return out.getvalue()
